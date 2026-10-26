import openpyxl

wb = openpyxl.load_workbook('milestone_order_form_2026.xlsx', data_only=True)

# Check Order Sheet for style code list (Table Lists sheet)
print("=== Table Lists sheet ===")
ws = wb['Table Lists']
for row in ws.iter_rows(values_only=True):
    if any(v is not None for v in row):
        print(row)

# Check if there's a style legend anywhere in the Price List sheet
print("\n=== Price List - first row (headers) ===")
ws = wb['Price List']
headers = next(ws.iter_rows(values_only=True))
print(headers)

# Check if there's a row 2 that has style names
print("\n=== Price List - rows 1-3 ===")
for i, row in enumerate(ws.iter_rows(values_only=True)):
    if i >= 3: break
    print(f"Row {i+1}:", row)

# Look for any sheet that maps style codes to names
print("\n=== All sheet names ===")
print(wb.sheetnames)

# Check Order Sheet row 1 for style name
print("\n=== Order Sheet - first 3 rows ===")
ws = wb['Order Sheet']
for i, row in enumerate(ws.iter_rows(values_only=True)):
    if i >= 3: break
    print(f"Row {i+1}:", row[:10])

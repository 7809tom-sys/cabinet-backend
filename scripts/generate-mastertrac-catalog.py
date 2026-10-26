#!/usr/bin/env python3
"""Export the approved MasterTrac worksheet to a customer-safe catalog.

Usage: python3 scripts/generate-mastertrac-catalog.py /path/to/approved.xlsm
Only MasterTrac rows and the customer List prices (E/F) are exported. The
private dealer Invoice column (D) is deliberately never read or serialized.
"""
import json
import re
import sys
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path

import openpyxl

SOURCE = Path(sys.argv[1]) if len(sys.argv) == 2 else None
if SOURCE is None or not SOURCE.is_file():
    raise SystemExit("Pass the path to the approved WoodTrac wholesale workbook.")
DEST = Path(__file__).resolve().parents[1] / "client/src/lib/mastertrac-catalog.json"

def decimal(value):
    if isinstance(value, bool) or not isinstance(value, (float, int)):
        raise ValueError(f"Expected cached numeric price, got {value!r}")
    return Decimal(str(value))

def dollars(value):
    return float(value.quantize(Decimal("0.01"), rounding=ROUND_HALF_UP))

def inches(value):
    if value is None:
        return None
    if isinstance(value, (float, int)):
        return float(value)
    match = re.match(r"^\s*(\d+(?:\.\d+)?|\.\d+)\s*(?:in|\"|”)?\s*$", str(value))
    return float(match.group(1)) if match else None

worksheet = openpyxl.load_workbook(SOURCE, read_only=True, data_only=True)["MasterTrac"]
products = []
seen = set()
for row_no, row in enumerate(worksheet.iter_rows(min_row=10, max_row=worksheet.max_row, values_only=True), 10):
    raw_sku = row[0]
    if raw_sku is None:
        continue
    if not isinstance(raw_sku, (int, float)) or int(raw_sku) != raw_sku:
        raise ValueError(f"MasterTrac row {row_no} has a nonnumeric material number: {raw_sku!r}")
    sku = str(int(raw_sku))
    if sku in seen:
        raise ValueError(f"Duplicate MasterTrac SKU {sku}")
    seen.add(sku)
    list_each, list_pack = decimal(row[4]), decimal(row[5])
    qty = row[12]
    if not isinstance(qty, (int, float)) or int(qty) != qty or qty < 1:
        raise ValueError(f"Invalid pack size for {sku}: {qty!r}")
    qty = int(qty)
    if list_each <= 0 or list_pack <= 0 or abs(list_pack - list_each * qty) > Decimal("0.01"):
        raise ValueError(f"List/pack mismatch for {sku} at row {row_no}")
    category, description, finish = (str(row[j]).strip() for j in (8, 6, 7))
    if not all((category, description, finish)):
        raise ValueError(f"Incomplete product metadata for {sku}")
    raw_length, raw_thickness, raw_depth = (inches(row[j]) for j in (9, 10, 11))
    panel = category in ("Panels (Hang)", "Panels (Floor)") and qty == 1
    shelf = category == "Shelves" and qty == 1 and "corner" not in description.lower() and "shoe" not in description.lower()
    placeable = (panel or shelf) and all(v and v > 0 for v in (raw_length, raw_thickness, raw_depth))
    if panel:
        width, depth, height = raw_thickness, raw_depth, raw_length
        placement = "wall" if category == "Panels (Hang)" else "tall"
        shape = "panel"
    elif shelf:
        width, depth, height = raw_length, raw_depth, raw_thickness
        placement = "wall"
        shape = "shelf"
    else:
        width = depth = height = placement = shape = None
    products.append({
        "sku": sku, "description": description, "finish": finish,
        "category": category, "listEachUsd": dollars(list_each),
        "listPackUsd": dollars(list_pack), "packQty": qty,
        "sourceRow": row_no, "lengthIn": raw_length,
        "thicknessIn": raw_thickness, "depthIn": raw_depth,
        "placeable": bool(placeable), "shape": shape,
        "placement": placement, "widthIn": width, "heightIn": height,
    })
if len(products) != 454 or len(seen) != 454:
    raise ValueError(f"Expected exactly 454 distinct MasterTrac SKUs, got {len(products)}")
DEST.write_text(json.dumps(products, ensure_ascii=False, separators=(",", ":")) + "\n")
print(f"Exported {len(products)} List-priced SKUs, {sum(p['placeable'] for p in products)} directly placeable; {DEST}")
